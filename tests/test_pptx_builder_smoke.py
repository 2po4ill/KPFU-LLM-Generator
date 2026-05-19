"""Smoke test: PPTX export with title, text-only, and image slide."""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from presentation.pptx_builder import build_pptx_bytes, prepare_slides_for_pptx

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app"
TEMPLATE = APP / "presentation" / "templates" / "kpfu_ru.pptx"
FIXTURE_PNG = (
    APP
    / "app"
    / "cache"
    / "books"
    / "presentation"
    / "d830d4c7ca94"
    / "11b4792011c7"
    / "mermaid"
    / "slide_5518586.png"
)


@pytest.fixture(autouse=True)
def _kpfu_template(monkeypatch):
    if not TEMPLATE.exists():
        pytest.skip(f"template missing: {TEMPLATE}")
    monkeypatch.setenv("PACKAGE_PPTX_TEMPLATE_PATH", str(TEMPLATE))
    monkeypatch.setenv("PACKAGE_PPTX_LAYOUT_TITLE", "TITLE")
    monkeypatch.setenv("PACKAGE_PPTX_LAYOUT_CONTENT", "TITLE_AND_BODY")
    monkeypatch.setenv("PACKAGE_PPTX_LAYOUT_CONTENT_IMAGE", "Два объекта")
    from core.config import Settings

    import core.config as cfg

    cfg.settings = Settings()


def test_pptx_smoke_three_slide_types(tmp_path: Path):
    if not FIXTURE_PNG.exists():
        pytest.skip(f"fixture png missing: {FIXTURE_PNG}")

    slides = [
        {
            "layout": "title",
            "title": "Уязвимости и угрозы",
            "subtitle": "Тема 2 · smoke test",
        },
        {
            "layout": "content",
            "title": "Классификация угроз",
            "bullets": [
                "Угроза — потенциальное нарушение безопасности.",
                "STRIDE группирует угрозы по типу воздействия.",
            ],
        },
        {
            "layout": "content",
            "title": "Модель STRIDE",
            "bullets": [
                "Spoofing, Tampering, Repudiation, Information disclosure, DoS, Elevation.",
            ],
            "image_caption": "Классификация угроз по модели STRIDE",
        },
    ]
    resolved = [None, FIXTURE_PNG]

    out = tmp_path / "smoke.pptx"
    out.write_bytes(
        build_pptx_bytes(
            slides,
            resolved,
            theme_title="Уязвимости и угрозы",
            wrap_with_title_slides=False,
        )
    )

    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    assert round(prs.slide_width / 914400, 3) == 10.0

    text_slide = prs.slides[1]
    assert text_slide.slide_layout.name == "TITLE_AND_BODY"
    title_ph = text_slide.shapes.title
    body_ph = next(
        ph for ph in text_slide.placeholders if ph.placeholder_format.idx == 1
    )
    assert title_ph.top < body_ph.top, "title must be above body"
    assert title_ph.text_frame.paragraphs[0].font.bold
    assert body_ph.text_frame.paragraphs[0].alignment != PP_ALIGN.CENTER

    image_slide = prs.slides[2]
    assert image_slide.slide_layout.name == "Два объекта"
    img_title = image_slide.shapes.title
    assert img_title.text_frame.paragraphs[0].font.bold
    assert img_title.top < 500000, "title near top of slide"

    pictures = [
        sh for sh in image_slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert pictures, "image on right side"
    pic = pictures[0]
    assert pic.left > prs.slide_width * 0.45

    body_left = next(
        ph for ph in image_slide.placeholders if ph.placeholder_format.idx == 1
    )
    assert body_left.width < prs.slide_width * 0.55, "bullets in left column"
    assert "STRIDE" in img_title.text_frame.text

    assert any(
        sh.has_text_frame and "Рис. 1" in sh.text_frame.text for sh in image_slide.shapes
    )


def test_prepare_slides_no_wrap():
    raw = [{"layout": "content", "title": "A", "bullets": ["b"]}]
    out = prepare_slides_for_pptx(raw, wrap_with_title_slides=False)
    assert len(out) == 1
    assert out[0]["title"] == "A"
