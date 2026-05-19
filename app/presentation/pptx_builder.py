"""
Build .pptx from resolved slides using layout spec or optional template .pptx.
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from presentation.layout_spec import (
    SlideLayoutSpec,
    TextBoxSpec,
    get_layout_spec,
    layout_names,
    resolve_template_path,
)
from presentation.sanitize import simplify_markdown_for_slide

logger = logging.getLogger(__name__)

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


_LAYOUT_ALIASES: Dict[str, List[str]] = {
    "title": ["титульный", "kpfu title", "title slide", "title only"],
    "title_and_body": ["только текст", "kpfu content", "title and content"],
    "два объекта": [
        "текст с изображением",
        "kpfu content image",
        "picture with caption",
        "рисунок с подписью",
    ],
    "титульный": ["kpfu title", "title slide", "title"],
    "только текст": ["title_and_body", "kpfu content", "title and content"],
    "текст с изображением": ["два объекта", "kpfu content image"],
}


def _find_layout(prs: "Presentation", name: str):
    primary = (name or "").strip().lower()
    candidates = [primary] + _LAYOUT_ALIASES.get(primary, [])
    for want in candidates:
        for layout in prs.slide_layouts:
            if (layout.name or "").strip().lower() == want:
                return layout
    return None


def prepare_slides_for_pptx(
    slides: List[Dict[str, Any]],
    *,
    theme_title: str = "",
    closing_subtitle: str = "",
    wrap_with_title_slides: bool = True,
) -> List[Dict[str, Any]]:
    """Title first, content in the middle, optional closing title slide."""
    items = [dict(s) for s in (slides or [])]
    if not items:
        items = [
            {
                "layout": "title",
                "title": theme_title or "Лекция",
                "subtitle": closing_subtitle,
            }
        ]
    if not wrap_with_title_slides:
        return items
    if str(items[0].get("layout") or "").lower() != "title":
        items.insert(
            0,
            {
                "layout": "title",
                "title": theme_title or "Лекция",
                "subtitle": closing_subtitle,
            },
        )
    opening = items[0]
    if str(items[-1].get("layout") or "").lower() != "title":
        items.append(
            {
                "layout": "title",
                "title": str(opening.get("title") or theme_title or "Лекция"),
                "subtitle": str(opening.get("subtitle") or closing_subtitle or ""),
            }
        )
    return items


def _blank_layout(prs: "Presentation"):
    for idx in (6, 5, len(prs.slide_layouts) - 1):
        if idx < len(prs.slide_layouts):
            return prs.slide_layouts[idx]
    return prs.slide_layouts[0]


def _strip_template_slides(prs: "Presentation") -> None:
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]


def _load_presentation() -> Tuple["Presentation", bool]:
    path = resolve_template_path()
    if path:
        logger.info("PPTX template: %s", path)
        prs = Presentation(str(path))
        _strip_template_slides(prs)
        return prs, True
    prs = Presentation()
    prs.slide_width = Inches(get_layout_spec().slide_width)
    prs.slide_height = Inches(get_layout_spec().slide_height)
    return prs, False


def _set_textbox(
    slide,
    spec: TextBoxSpec,
    text: str,
    *,
    font_pt: int,
    bold: bool = False,
    align_center: bool = False,
) -> None:
    box = slide.shapes.add_textbox(
        Inches(spec.left),
        Inches(spec.top),
        Inches(spec.width),
        Inches(spec.height),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_pt)
    p.font.bold = bold
    if align_center:
        p.alignment = PP_ALIGN.CENTER


def _add_bullets_to_box(
    slide,
    spec: TextBoxSpec,
    bullets: List[str],
    *,
    font_pt: int,
    spec_limits: SlideLayoutSpec,
    align_center: bool = False,
) -> None:
    box = slide.shapes.add_textbox(
        Inches(spec.left),
        Inches(spec.top),
        Inches(spec.width),
        Inches(spec.height),
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    trimmed = [
        simplify_markdown_for_slide(b, max_len=spec_limits.max_bullet_chars)
        for b in bullets[: spec_limits.max_bullets]
    ]
    if not trimmed:
        trimmed = ["—"]
    for i, bullet in enumerate(trimmed):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(font_pt)
        if align_center:
            p.alignment = PP_ALIGN.CENTER


def _image_size_for_box(img_path: Path, box: TextBoxSpec) -> Tuple[float, float]:
    max_w = box.width
    max_h = box.height
    if HAS_PIL:
        with Image.open(img_path) as im:
            w, h = im.size
        if w <= 0 or h <= 0:
            return max_w, max_h
        scale = min(max_w / w, max_h / h)
        return w * scale, h * scale
    return max_w * 0.92, max_h * 0.92


def format_figure_caption(slide_def: Dict[str, Any]) -> str:
    num = slide_def.get("figure_number")
    note = str(slide_def.get("image_caption") or slide_def.get("image_note") or "").strip()
    if num is None:
        return note
    prefix = f"Рис. {int(num)}"
    if note:
        return f"{prefix} - {note}"
    return prefix


def _place_image_caption(slide, text: str, box: TextBoxSpec) -> None:
    if not text:
        return
    _set_textbox(slide, box, text, font_pt=11, align_center=True)


def _place_image(slide, img_path: Path, box: TextBoxSpec) -> None:
    w_in, h_in = _image_size_for_box(img_path, box)
    left = box.left + (box.width - w_in) / 2.0
    top = box.top + (box.height - h_in) / 2.0
    slide.shapes.add_picture(
        str(img_path),
        Inches(left),
        Inches(top),
        width=Inches(w_in),
        height=Inches(h_in),
    )


def _fill_placeholder_shape(
    shape,
    text: str,
    font_pt: int,
    *,
    bold: bool = False,
    align_center: bool = False,
) -> None:
    if not shape or not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_pt)
    p.font.bold = bold
    if align_center:
        p.alignment = PP_ALIGN.CENTER


def _fill_bullets_placeholder(
    body_ph,
    bullets: List[str],
    limits: SlideLayoutSpec,
    *,
    align_center: bool = False,
) -> None:
    tf = body_ph.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()
    trimmed = [
        simplify_markdown_for_slide(b, max_len=limits.max_bullet_chars)
        for b in bullets[: limits.max_bullets]
    ]
    if not trimmed:
        trimmed = ["—"]
    for i, bullet in enumerate(trimmed):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(limits.bullet_font_pt)
        if align_center:
            p.alignment = PP_ALIGN.CENTER


def _placeholder_by_idx(slide, idx: int):
    for ph in slide.placeholders:
        if getattr(ph.placeholder_format, "idx", None) == idx:
            return ph
    return None


def _fill_title_placeholder(slide, title: str, limits: SlideLayoutSpec) -> None:
    if not slide.shapes.title:
        return
    _fill_placeholder_shape(
        slide.shapes.title,
        simplify_markdown_for_slide(title, max_len=limits.title_max_chars),
        limits.content_title_font_pt,
        bold=True,
    )


def _try_template_title(slide, title: str, subtitle: str, limits: SlideLayoutSpec) -> bool:
    try:
        if slide.shapes.title:
            _fill_placeholder_shape(
                slide.shapes.title,
                simplify_markdown_for_slide(title, max_len=limits.title_max_chars),
                limits.title_font_pt,
                bold=True,
            )
        sub_ph = _placeholder_by_idx(slide, 1)
        if sub_ph is not None:
            _fill_placeholder_shape(
                sub_ph,
                simplify_markdown_for_slide(subtitle, max_len=limits.subtitle_max_chars),
                20,
            )
        return True
    except Exception:
        return False


def _try_template_content_text(
    slide,
    title: str,
    bullets: List[str],
    limits: SlideLayoutSpec,
) -> bool:
    """TITLE_AND_BODY: bold title on top, left-aligned bullets below."""
    try:
        _fill_title_placeholder(slide, title, limits)
        body_ph = _placeholder_by_idx(slide, 1)
        if body_ph is None:
            return False
        _fill_bullets_placeholder(body_ph, bullets, limits, align_center=False)
        return True
    except Exception as e:
        logger.debug("template text content: %s", e)
        return False


def _try_template_content_split(
    slide,
    title: str,
    bullets: List[str],
    limits: SlideLayoutSpec,
) -> bool:
    """Два объекта: bold title on top, bullets in left column."""
    try:
        _fill_title_placeholder(slide, title, limits)
        body_ph = _placeholder_by_idx(slide, 1)
        if body_ph is None:
            return False
        _fill_bullets_placeholder(body_ph, bullets, limits, align_center=False)
        return True
    except Exception as e:
        logger.debug("template split content: %s", e)
        return False


def _apply_resolved_figure_numbers(
    slides: List[Dict[str, Any]],
    resolved_images: List[Optional[Path]],
) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    img_i = 0
    fig = 0
    for slide in slides:
        s = dict(slide)
        if str(s.get("layout") or "content").lower() == "title":
            s.pop("figure_number", None)
            out.append(s)
            continue
        path = resolved_images[img_i] if img_i < len(resolved_images) else None
        img_i += 1
        if path and Path(path).exists():
            fig += 1
            s["figure_number"] = fig
        elif not s.get("image_url") and not s.get("image_path"):
            s.pop("figure_number", None)
        out.append(s)
    return out


def build_pptx_bytes(
    slides: List[Dict[str, Any]],
    resolved_images: List[Optional[Path]],
    *,
    theme_title: str = "",
    wrap_with_title_slides: bool = True,
) -> bytes:
    if not HAS_PPTX:
        raise RuntimeError("python-pptx is not installed")

    from presentation.slides_export import assign_figure_numbers

    limits = get_layout_spec()
    wrapped = prepare_slides_for_pptx(
        slides,
        theme_title=theme_title,
        wrap_with_title_slides=wrap_with_title_slides,
    )
    slides = _apply_resolved_figure_numbers(
        assign_figure_numbers(wrapped),
        resolved_images,
    )
    prs, has_template = _load_presentation()
    title_layout_name, content_layout_name, content_image_layout_name = layout_names()

    title_layout = _find_layout(prs, title_layout_name) if has_template else None
    content_layout = _find_layout(prs, content_layout_name) if has_template else None
    content_image_layout = _find_layout(prs, content_image_layout_name) if has_template else None
    blank = _blank_layout(prs)

    if has_template and not title_layout:
        logger.warning("Template layout not found: %s", title_layout_name)
    if has_template and not content_layout:
        logger.warning("Template layout not found: %s", content_layout_name)
    if has_template and not content_image_layout:
        logger.warning("Template layout not found: %s", content_image_layout_name)

    img_idx = 0
    for slide_def in slides:
        layout_kind = str(slide_def.get("layout") or "content").lower()
        notes = str(slide_def.get("speaker_notes") or "")
        if slide_def.get("_image_note"):
            notes = "\n".join(filter(None, [notes, str(slide_def.get("_image_note"))]))

        if layout_kind == "title":
            title = simplify_markdown_for_slide(
                str(slide_def.get("title") or theme_title or "Лекция"),
                max_len=limits.title_max_chars,
            )
            subtitle = simplify_markdown_for_slide(
                str(slide_def.get("subtitle") or ""),
                max_len=limits.subtitle_max_chars,
            )
            layout = title_layout or blank
            slide = prs.slides.add_slide(layout)
            if title_layout and not _try_template_title(slide, title, subtitle, limits):
                _set_textbox(
                    slide,
                    limits.title_slide_title,
                    title,
                    font_pt=limits.title_font_pt,
                    bold=True,
                    align_center=True,
                )
                _set_textbox(
                    slide,
                    limits.title_slide_subtitle,
                    subtitle,
                    font_pt=20,
                    align_center=True,
                )
            elif not title_layout:
                _set_textbox(
                    slide,
                    limits.title_slide_title,
                    title,
                    font_pt=limits.title_font_pt,
                    bold=True,
                    align_center=True,
                )
                _set_textbox(
                    slide,
                    limits.title_slide_subtitle,
                    subtitle,
                    font_pt=20,
                    align_center=True,
                )
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            continue

        bullets = [
            str(b)
            for b in (slide_def.get("bullets") or [])
            if str(b).strip() and not str(b).strip().startswith("<!--")
        ]
        title = simplify_markdown_for_slide(
            str(slide_def.get("title") or "Слайд"),
            max_len=limits.title_max_chars,
        )
        img_path = resolved_images[img_idx] if img_idx < len(resolved_images) else None
        img_idx += 1
        has_image = bool(img_path and Path(img_path).exists())
        caption = format_figure_caption(slide_def)

        use_image_layout = has_image and content_image_layout is not None
        layout = content_image_layout if use_image_layout else content_layout
        if not layout:
            layout = blank
        slide = prs.slides.add_slide(layout)

        if use_image_layout:
            used = (
                _try_template_content_split(slide, title, bullets, limits)
                if content_image_layout
                else False
            )
            if not used:
                _set_textbox(
                    slide,
                    limits.content_title,
                    title,
                    font_pt=limits.content_title_font_pt,
                    bold=True,
                )
                _add_bullets_to_box(
                    slide,
                    limits.content_body,
                    bullets,
                    font_pt=limits.bullet_font_pt,
                    spec_limits=limits,
                    align_center=False,
                )
            _place_image(slide, Path(img_path), limits.content_image)
            if caption:
                _place_image_caption(slide, caption, limits.content_image_caption)
        else:
            used = (
                _try_template_content_text(slide, title, bullets, limits)
                if content_layout
                else False
            )
            if not used:
                _set_textbox(
                    slide,
                    limits.content_title,
                    title,
                    font_pt=limits.content_title_font_pt,
                    bold=True,
                )
                _add_bullets_to_box(
                    slide,
                    limits.content_body_full,
                    bullets,
                    font_pt=limits.bullet_font_pt,
                    spec_limits=limits,
                    align_center=False,
                )

        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_pptx_file(
    slides: List[Dict[str, Any]],
    resolved_images: List[Optional[Path]],
    out_path: Path,
    *,
    theme_title: str = "",
    wrap_with_title_slides: bool = True,
) -> Path:
    data = build_pptx_bytes(
        slides,
        resolved_images,
        theme_title=theme_title,
        wrap_with_title_slides=wrap_with_title_slides,
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(data)
    return out_path
